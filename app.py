import json
import os
import tempfile
import uuid
import threading
import itertools
import logging
import traceback
from io import BytesIO
from typing import Dict, List, Any
from datetime import datetime
from flask import Flask, render_template, request, redirect, url_for, flash
from flask_login import LoginManager, current_user, login_required, login_user, logout_user
from email_validator import validate_email, EmailNotValidError
from config import Config
from models import db, User, Store, ProductSnapshot, ActionSheet
from ozon_client import OzonClient, OzonClientError
from utils import PACKAGING_COST, PROMOTION_RATE, compute_financials, prepare_table_data
import smtplib
from email.message import EmailMessage
from sqlalchemy import inspect, text
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


app = Flask(__name__)
app.config.from_object(Config)
db.init_app(app)

# Базовый лог в файл и консоль, чтобы видеть причину сбоев фоновых задач
logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(name)s: %(message)s",
    handlers=[
        logging.StreamHandler(),
        logging.FileHandler("fetch_jobs.log", encoding="utf-8"),
    ],
)
logger = logging.getLogger("fetch_jobs")

STATUS_LABELS = {
    "awaiting_packaging": "Ожидает упаковки",
    "awaiting_deliver": "Ожидает отгрузки",
    "awaiting_registration": "Ожидает регистрации",
    "awaiting_approve": "Ожидает подтверждения",
    "awaiting_verification": "Создано",
    "acceptance_in_progress": "Идёт приёмка",
    "delivering": "Доставляется",
    "delivered": "Доставлено",
    "driver_pickup": "У водителя",
    "not_accepted": "Не принят",
    "arbitration": "Арбитраж",
    "client_arbitration": "Клиентский арбитраж",
    "sent_by_seller": "Отправлено продавцом",
    "cancelled": "Отменено",
    "cancelled_from_split_pending": "Отменено",
    "posting_created": "Создано",
}

fetch_jobs: Dict[str, Dict] = {}
fetch_lock = threading.Lock()

login_manager = LoginManager(app)
login_manager.login_view = "login"


@login_manager.user_loader
def load_user(user_id):
    try:
        return db.session.get(User, int(user_id))
    except Exception:
        return None


def ensure_admin_user():
    admin = User.query.filter_by(email=Config.ADMIN_EMAIL).first()
    if not admin:
        admin = User(email=Config.ADMIN_EMAIL, is_admin=True)
        admin.set_password(Config.ADMIN_PASSWORD)
        db.session.add(admin)
        db.session.commit()


def ensure_schema():
    inspector = inspect(db.engine)
    if not inspector.has_table("user"):
        return
    column_names = {col["name"] for col in inspector.get_columns("user")}
    with db.engine.begin() as conn:
        if "is_admin" not in column_names:
            conn.execute(text("ALTER TABLE user ADD COLUMN is_admin BOOLEAN DEFAULT 0"))
        if "is_limited" not in column_names:
            conn.execute(text("ALTER TABLE user ADD COLUMN is_limited BOOLEAN DEFAULT 0"))


@app.before_request
def create_tables():
    db.create_all()
    ensure_schema()
    ensure_admin_user()


def send_email(to_email: str, subject: str, body: str):
    if not Config.MAIL_SERVER:
        print(f"Email fallback (set MAIL_SERVER to send real email):\nTo: {to_email}\nSubject: {subject}\n{body}")
        return
    msg = EmailMessage()
    msg["Subject"] = subject
    msg["From"] = Config.MAIL_DEFAULT_SENDER
    msg["To"] = to_email
    msg.set_content(body)

    with smtplib.SMTP(Config.MAIL_SERVER, Config.MAIL_PORT) as server:
        if Config.MAIL_USE_TLS:
            server.starttls()
        if Config.MAIL_USERNAME and Config.MAIL_PASSWORD:
            server.login(Config.MAIL_USERNAME, Config.MAIL_PASSWORD)
        server.send_message(msg)


def _parse_dt(value: str, fallback: datetime) -> datetime:
    if not value:
        return fallback
    try:
        if value.endswith("Z"):
            return datetime.fromisoformat(value.replace("Z", "+00:00"))
        return datetime.fromisoformat(value)
    except Exception:
        return fallback


def _iso_utc(dt_value: datetime) -> str:
    return dt_value.replace(microsecond=0).isoformat() + "Z"


def _parse_action_sheet(file_storage) -> Dict[str, Any]:
    # читаем файл дважды: с формулами и с вычисленными значениями
    content = file_storage.read()
    wb_values = load_workbook(BytesIO(content), data_only=True)
    wb_formulas = load_workbook(BytesIO(content), data_only=False)

    # Ищем вкладку «Товары и цены» (или похожее название), чтобы не зависеть от порядка листов
    sheet_values = wb_values.active
    sheet_formulas = wb_formulas[sheet_values.title]
    for name in wb_values.sheetnames:
        lower = name.lower()
        if "товар" in lower and "цен" in lower:
            sheet_values = wb_values[name]
            sheet_formulas = wb_formulas[name]
            break

    target_fields = [
        {"labels": ["артикул"], "key": "article", "label": "Артикул", "display": True},
        {"labels": ["название"], "key": "name", "label": "Название", "display": True},
        {"labels": ["цена до скидки, rub", "цена до скидки"], "key": "old_price", "label": "Цена до скидки, RUB", "display": True},
        {"labels": ["ваша цена, rub", "ваша цена"], "key": "your_price", "label": "Ваша цена, RUB", "display": True},
        {"labels": ["текущая цена, rub", "текущая цена"], "key": "current_price", "label": "Текущая цена, RUB", "display": True},
        {"labels": ["минимальная цена, rub", "минимальная цена"], "key": "min_price", "label": "Минимальная цена, RUB", "display": True},
        {
            "labels": [
                "участие товара в акции",
                "участие товара в глобальной акции",
                "участие товара в глобальной акции (можно редактировать)",
            ],
            "key": "participate",
            "label": "Участие товара в акции",
            "display": True,
        },
        {"labels": ["итоговая цена по акции, rub", "итоговая цена по акции"], "key": "action_price", "label": "Итоговая цена по акции, RUB", "display": True},
        {"labels": ["количество товаров в акции, шт"], "key": "qty", "label": "Количество товаров в акции, шт", "display": True},
        {"labels": ["акционный бустинг в поиске"], "key": "current_boost", "label": "Акционный бустинг в поиске", "display": True},
        {"labels": ["цена для минимального акционного бустинга, rub"], "key": "price_min_elastic", "label": "Цена для минимального акционного бустинга, RUB", "display": True},
        {"labels": ["цена для максимального акционного бустинга, rub"], "key": "price_max_elastic", "label": "Цена для максимального акционного бустинга, RUB", "display": True},
        {"labels": ["ozon id", "ozon_id", "product id", "product_id", "ozonid", "ozon id (обязательное поле)", "ozon id (обязательное поле)"], "key": "product_id", "label": "Ozon ID", "display": False},
    ]

    def _as_bool(value: Any) -> bool:
        if value is None:
            return False
        if isinstance(value, (int, float)):
            return value != 0
        text = str(value).strip().lower()
        return text in {"1", "true", "да", "yes", "y", "on", "x"}

    def _normalize(value: Any) -> str:
        return str(value).strip().lower().replace("  ", " ") if value is not None else ""

    # Найдём строку заголовков, где присутствует колонка «Артикул»
    header_row_index = None
    for idx, row in enumerate(sheet_values.iter_rows(values_only=True), start=1):
        if any(_normalize(cell) == "артикул" for cell in row):
            header_row_index = idx
            break

    if header_row_index is None:
        raise ValueError("Не удалось найти строку с заголовками (Артикул)")

    raw_header_row = next(
        sheet_values.iter_rows(min_row=header_row_index, max_row=header_row_index, values_only=True)
    )
    header_map: Dict[int, Dict[str, Any]] = {}
    headers: List[Dict[str, str]] = []

    for idx, cell in enumerate(raw_header_row):
        norm_cell = _normalize(cell)
        for field in target_fields:
            if norm_cell in [_normalize(lbl) for lbl in field["labels"]]:
                header_map[idx] = field
                if field["display"]:
                    headers.append({"label": field["label"], "key": field["key"]})
                break

    if not header_map:
        raise ValueError("В заголовке не найдено ни одной нужной колонки")

    rows: List[Dict[str, Any]] = []

    max_columns = sheet_values.max_column

    def _extract_product_id(mapped: Dict[str, Any]):
        candidates = [
            mapped.get("product_id"),
            mapped.get("ozon_id"),
            mapped.get("ozonid"),
            mapped.get("id"),
        ]
        for key, value in list(mapped.items()):
            if value is None:
                continue
            if "ozon" in key and "id" in key and value not in candidates:
                candidates.append(value)

        for candidate in candidates:
            if candidate in (None, ""):
                continue
            try:
                if isinstance(candidate, float) and candidate.is_integer():
                    candidate = int(candidate)
                return candidate
            except Exception:
                return candidate
        return None

    def _cell_letter(col_idx: int) -> str:
        return get_column_letter(col_idx + 1)

    # определяем ключевые колонки для последующей пересборки формул
    action_price_col_idx = next((i for i, f in header_map.items() if f["key"] == "action_price"), None)
    participate_col_idx = next((i for i, f in header_map.items() if f["key"] == "participate"), None)
    boost_col_idx = next((i for i, f in header_map.items() if f["key"] == "current_boost"), None)

    for relative_idx, row in enumerate(
        sheet_values.iter_rows(min_row=header_row_index + 1, values_only=True)
    ):
        if all(v is None for v in row):
            continue
        mapped: Dict[str, Any] = {}
        raw_map: Dict[str, Any] = {}
        values_row: List[Any] = [None] * max_columns
        for idx in range(max_columns):
            value = row[idx] if idx < len(row) else None
            if idx in header_map:
                field = header_map[idx]
                if field["key"] == "participate":
                    value = 1 if _as_bool(value) else 0
                mapped[field["key"]] = value
            values_row[idx] = value
            norm_key = _normalize(raw_header_row[idx]) if idx < len(raw_header_row) else f"column_{idx}"
            raw_map[norm_key] = value
        product_id = _extract_product_id(raw_map)
        if product_id in (None, ""):
            continue
        mapped["product_id"] = product_id

        boost_formula = None
        if boost_col_idx is not None:
            col_letter = _cell_letter(boost_col_idx)
            excel_row = header_row_index + 1 + relative_idx
            cell_key = f"{col_letter}{excel_row}"
            cell = sheet_formulas[cell_key]
            boost_formula = cell.value if cell.data_type == "f" else cell.value
            if hasattr(cell, "_value") and cell.data_type == "f":
                boost_formula = cell._value

        rows.append(
            {
                "data": mapped,
                "meta": {
                    "values": values_row,
                    "boost_formula": boost_formula or "",
                    "boost_col": boost_col_idx,
                    "action_price_col": action_price_col_idx,
                    "participate_col": participate_col_idx,
                },
            }
        )

    return {"headers": headers, "rows": rows}


def _interpolate_boost(price: float, row: Dict[str, Any]) -> float:
    try:
        p_min = float(row.get("price_min_elastic") or 0)
        p_max = float(row.get("price_max_elastic") or 0)
        b_min = float(row.get("min_boost") or 0)
        b_max = float(row.get("max_boost") or 0)
        if p_min == 0 and p_max == 0:
            return float(row.get("current_boost") or 0)
        if p_min == p_max:
            return b_max or b_min or float(row.get("current_boost") or 0)
        # чем ниже цена, тем выше буст — линейная интерполяция в пределах эластики
        clamped = max(min(price, p_min), p_max) if p_min > p_max else max(min(price, p_max), p_min)
        span = (p_min - p_max) if p_min != p_max else 1
        ratio = (p_min - clamped) / span
        return b_min + (b_max - b_min) * ratio
    except Exception:
        return float(row.get("current_boost") or 0)


@app.route("/register", methods=["GET", "POST"])
def register():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        try:
            validate_email(email)
        except EmailNotValidError:
            flash("Введите корректный email", "danger")
            return render_template("auth/register.html")
        if len(password) < 6:
            flash("Пароль должен быть не меньше 6 символов", "danger")
            return render_template("auth/register.html")
        if User.query.filter_by(email=email).first():
            flash("Пользователь уже существует", "warning")
            return render_template("auth/register.html")
        user = User(email=email)
        user.set_password(password)
        db.session.add(user)
        db.session.commit()
        login_user(user)
        return redirect(url_for("dashboard"))
    return render_template("auth/register.html")


@app.route("/login", methods=["GET", "POST"])
def login():
    if current_user.is_authenticated:
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        password = request.form.get("password", "")
        user = User.query.filter_by(email=email).first()
        if user and user.check_password(password):
            login_user(user)
            return redirect(url_for("dashboard"))
        flash("Неверные учетные данные", "danger")
    return render_template("auth/login.html")


@app.route("/logout")
@login_required
def logout():
    logout_user()
    return redirect(url_for("login"))


@app.route("/admin/users", methods=["GET", "POST"])
@login_required
def admin_users():
    if not current_user.is_admin:
        flash("Недостаточно прав", "danger")
        return redirect(url_for("dashboard"))
    if request.method == "POST":
        target_id = request.form.get("user_id", type=int)
        toggle_field = request.form.get("field")
        user = User.query.get_or_404(target_id)
        if toggle_field == "limited":
            user.is_limited = not user.is_limited
            db.session.commit()
            flash("Ограничения обновлены", "info")
        elif toggle_field == "admin" and user.id != current_user.id:
            user.is_admin = not user.is_admin
            db.session.commit()
            flash("Роль обновлена", "success")
        return redirect(url_for("admin_users"))
    users = User.query.order_by(User.created_at.desc()).all()
    return render_template("admin_users.html", users=users)


@app.route("/reset", methods=["GET", "POST"])
def reset_request():
    if request.method == "POST":
        email = request.form.get("email", "").strip().lower()
        user = User.query.filter_by(email=email).first()
        if user:
            token = user.generate_reset_token()
            link = url_for('reset_token', token=token, _external=True)
            send_email(email, f"Сброс пароля {Config.APP_NAME}", f"Чтобы сбросить пароль, перейдите по ссылке: {link}")
            flash("Проверьте почту для восстановления пароля", "info")
        else:
            flash("Пользователь не найден", "warning")
    return render_template("auth/reset_request.html")


@app.route("/reset/<token>", methods=["GET", "POST"])
def reset_token(token):
    user = User.verify_reset_token(token)
    if not user:
        flash("Недействительная или истекшая ссылка", "danger")
        return redirect(url_for("reset_request"))
    if request.method == "POST":
        password = request.form.get("password", "")
        if len(password) < 6:
            flash("Пароль должен быть не меньше 6 символов", "danger")
            return render_template("auth/reset_token.html")
        user.set_password(password)
        db.session.commit()
        flash("Пароль обновлен", "success")
        return redirect(url_for("login"))
    return render_template("auth/reset_token.html")


@app.route("/")
@login_required
def dashboard():
    stores_query = Store.query
    snapshot_query = ProductSnapshot.query
    if not current_user.is_admin:
        stores_query = stores_query.filter_by(user_id=current_user.id)
        snapshot_query = snapshot_query.filter_by(user_id=current_user.id)
    stores_count = stores_query.count()
    last_snapshot = snapshot_query.order_by(ProductSnapshot.created_at.desc()).first()
    total_products = 0
    if last_snapshot and isinstance(last_snapshot.data, list):
        total_products = len(last_snapshot.data)
    return render_template("dashboard.html", stores_count=stores_count, total_products=total_products)


@app.route("/stores", methods=["GET", "POST"])
@login_required
def stores():
    users = []
    if current_user.is_admin:
        users = User.query.order_by(User.email).all()
    if request.method == "POST":
        if current_user.is_limited and not current_user.is_admin:
            flash("Ваш доступ к изменениям ограничен администратором", "warning")
            return redirect(url_for("stores"))
        name = request.form.get("name", "").strip()
        client_id = request.form.get("client_id", "").strip()
        api_key = request.form.get("api_key", "").strip()
        if not all([name, client_id, api_key]):
            flash("Заполните все поля", "danger")
        else:
            store_id = request.form.get("store_id")
            if store_id:
                store = Store.query.filter_by(id=store_id).first() if current_user.is_admin else Store.query.filter_by(id=store_id, user_id=current_user.id).first()
                if store:
                    store.name = name
                    store.client_id = client_id
                    store.api_key = api_key
                    if current_user.is_admin:
                        store.user_id = int(request.form.get("user_id") or store.user_id)
                    flash("Магазин обновлен", "success")
            else:
                owner = current_user
                if current_user.is_admin:
                    selected_user_id = request.form.get("user_id", type=int)
                    owner = User.query.get(selected_user_id) or current_user
                store = Store(name=name, client_id=client_id, api_key=api_key, owner=owner)
                db.session.add(store)
                flash("Магазин добавлен", "success")
            db.session.commit()
            return redirect(url_for("stores"))
    stores_list = Store.query.all() if current_user.is_admin else Store.query.filter_by(user_id=current_user.id).all()
    return render_template("stores.html", stores=stores_list, users=users)


@app.route("/stores/<int:store_id>/delete", methods=["POST"])
@login_required
def delete_store(store_id):
    store = Store.query.filter_by(id=store_id).first_or_404()
    if not current_user.is_admin and store.user_id != current_user.id:
        flash("Нет прав на удаление магазина", "danger")
        return redirect(url_for("stores"))
    db.session.delete(store)
    db.session.commit()
    flash("Магазин удален", "info")
    return redirect(url_for("stores"))


@app.route("/orders", methods=["GET", "POST"])
@login_required
def orders():
    stores = Store.query.all() if current_user.is_admin else Store.query.filter_by(user_id=current_user.id).all()
    packaging_override = request.values.get("packaging", type=float)
    promotion_override = request.values.get("promotion", type=float)
    promotion_rate = promotion_override / 100 if promotion_override is not None else None
    selected_store_ids_raw = request.values.getlist("store_ids")
    selected_store_ids = [int(s) for s in selected_store_ids_raw if s.isdigit()]
    include_all = "all" in selected_store_ids_raw

    today = datetime.utcnow()
    default_since = today.replace(hour=0, minute=0, second=0, microsecond=0)
    since_value = request.values.get("since") or _iso_utc(default_since)
    to_value = request.values.get("to") or _iso_utc(today)

    table_data = []
    product_lookup = {}

    allowed_store_ids = {store.id for store in stores}
    target_store_ids = allowed_store_ids if include_all or not selected_store_ids else set(selected_store_ids)

    for store_id in target_store_ids:
        latest_snapshot = (
            ProductSnapshot.query.filter_by(store_id=store_id)
            .order_by(ProductSnapshot.created_at.desc())
            .first()
        )
        if latest_snapshot:
            prepared = prepare_table_data(latest_snapshot.data, packaging_override, promotion_rate)
            for item in prepared:
                product_lookup[item.get("offer_id")] = item

    if request.method == "POST" and request.form.get("action") == "fetch":
        fetch_ids_raw = request.form.getlist("store_ids")
        fetch_all = "all" in fetch_ids_raw
        fetch_ids = [int(s) for s in fetch_ids_raw if s.isdigit()]
        target_ids = allowed_store_ids if fetch_all or not fetch_ids else set(fetch_ids)

        if not target_ids:
            flash("Выберите магазин", "warning")
            return redirect(url_for("orders"))

        since_dt = _parse_dt(request.form.get("since"), default_since)
        to_dt = _parse_dt(request.form.get("to"), today)
        since_iso = _iso_utc(since_dt)
        to_iso = _iso_utc(to_dt)
        since_value = since_iso
        to_value = to_iso

        postings = []
        try:
            for store in stores:
                if store.id not in target_ids:
                    continue
                if not current_user.is_admin and store.user_id != current_user.id:
                    flash("Нет доступа к одному из магазинов", "danger")
                    return redirect(url_for("orders"))
                client = OzonClient(store.client_id, store.api_key)
                postings.extend(client.fetch_postings(since_iso, to_iso))
            table_data = _prepare_postings(postings, product_lookup, packaging_override, promotion_rate)
            flash(f"Загружено отправлений: {len(table_data)}", "success")
        except OzonClientError as exc:
            flash(str(exc), "danger")
        except Exception as exc:
            flash(f"Не удалось получить отправления: {exc}", "danger")

    status_groups: Dict[str, List[Dict[str, Any]]] = {}
    for order in table_data:
        key = order.get("status_key") or "прочее"
        status_groups.setdefault(key, []).append(order)

    status_order = [
        "awaiting_packaging",
        "awaiting_deliver",
        "awaiting_registration",
        "awaiting_approve",
        "delivering",
        "delivered",
        "cancelled",
    ]
    ordered_keys = [key for key in status_order if key in status_groups]
    ordered_keys.extend(sorted(k for k in status_groups.keys() if k not in ordered_keys))
    status_tabs = [
        {
            "key": key,
            "label": STATUS_LABELS.get(key, key or "Статус"),
            "items": status_groups[key],
        }
        for key in ordered_keys
    ]

    active_orders = [o for o in table_data if not o.get("is_cancelled")]
    summary = {
        "amount": sum(o.get("amount", 0) for o in active_orders),
        "revenue": sum(o.get("revenue", 0) for o in active_orders),
        "expenses": sum(o.get("expenses", 0) for o in active_orders),
        "margin": sum(o.get("margin", 0) for o in active_orders),
        "cost": sum(o.get("cost", 0) for o in active_orders),
    }
    summary["markup"] = (summary["margin"] / summary["cost"] * 100) if summary["cost"] else 0

    return render_template(
        "orders.html",
        stores=stores,
        selected_store_ids=list(target_store_ids if table_data or selected_store_ids_raw else set()),
        since_value=since_value,
        to_value=to_value,
        table_data=table_data,
        status_tabs=status_tabs,
        summary=summary,
        packaging_value=packaging_override if packaging_override is not None else PACKAGING_COST,
        promotion_value=promotion_override if promotion_override is not None else PROMOTION_RATE * 100,
    )


@app.route("/actions", methods=["GET", "POST"])
@login_required
def actions_page():
    stores = Store.query.all() if current_user.is_admin else Store.query.filter_by(user_id=current_user.id).all()
    selected_store_id = request.values.get("store_id", type=int)
    actions_result = []
    selected_store = None
    last_sheet = None
    sheet_headers: List[Dict[str, Any]] = []
    sheet_rows: List[Dict[str, Any]] = []

    upload_file = request.files.get("sheet") if request.method == "POST" else None
    upload_target_store = request.form.get("store_id", type=int) if upload_file else None

    if selected_store_id:
        selected_store = next((s for s in stores if s.id == selected_store_id), None)
    elif stores:
        selected_store = stores[0]

    if upload_file and upload_target_store:
        target_store = next((s for s in stores if s.id == upload_target_store), None)
        if target_store and (current_user.is_admin or target_store.user_id == current_user.id):
            try:
                rows = _parse_action_sheet(upload_file)
                row_count = len(rows.get("rows", [])) if isinstance(rows, dict) else 0
                sheet = ActionSheet(
                    store_id=target_store.id,
                    user_id=current_user.id,
                    filename=upload_file.filename or "actions.xlsx",
                    data=rows,
                )
                db.session.add(sheet)
                db.session.commit()
                flash(
                    f"Файл {sheet.filename} загружен ({row_count} строк)",
                    "success",
                )
                selected_store = target_store
            except Exception as exc:
                db.session.rollback()
                flash(f"Не удалось обработать файл: {exc}", "danger")

    allowed_product_ids: set[str] = set()
    skipped_by_store = 0
    no_snapshot = False

    if selected_store and (current_user.is_admin or selected_store.user_id == current_user.id):
        last_sheet = (
            ActionSheet.query.filter_by(store_id=selected_store.id)
            .order_by(ActionSheet.uploaded_at.desc())
            .first()
        )
        if last_sheet:
            raw_data = last_sheet.data or {}
            if isinstance(raw_data, dict):
                sheet_headers = raw_data.get("headers", []) or []
                sheet_rows = raw_data.get("rows", []) or []
            elif isinstance(raw_data, list):
                # совместимость со старыми загрузками
                sheet_rows = [{"data": row, "meta": {}} for row in raw_data]
                if raw_data:
                    first_keys = list(raw_data[0].keys())
                    sheet_headers = [{"label": key, "key": key} for key in first_keys]

        latest_products = (
            ProductSnapshot.query.filter_by(store_id=selected_store.id)
            .order_by(ProductSnapshot.created_at.desc())
            .first()
        )
        if latest_products and isinstance(latest_products.data, list):
            for item in latest_products.data:
                pid_val = item.get("product_id") if isinstance(item, dict) else None
                if pid_val in (None, ""):
                    continue
                try:
                    if isinstance(pid_val, float) and pid_val.is_integer():
                        pid_val = int(pid_val)
                    allowed_product_ids.add(str(pid_val))
                except Exception:
                    allowed_product_ids.add(str(pid_val))
        elif sheet_rows:
            no_snapshot = True
            flash(
                "Сначала обновите товары для выбранного магазина, чтобы сверить Ozon ID из файла.",
                "warning",
            )

        client = OzonClient(selected_store.client_id, selected_store.api_key)
        try:
            actions_resp = client.list_actions()
            actions_list = actions_resp.get("result", []) if isinstance(actions_resp, dict) else []
            for action in actions_list:
                action_id = action.get("id")
                if not action_id:
                    continue
                candidates: List[Dict[str, Any]] = []
                if sheet_rows and not no_snapshot:
                    for row in sheet_rows:
                        data = row.get("data") if isinstance(row, dict) else row
                        meta = row.get("meta") if isinstance(row, dict) else {}
                        if not isinstance(data, dict):
                            continue
                        pid_val = data.get("product_id")
                        if pid_val not in (None, ""):
                            pid_str = str(int(pid_val)) if isinstance(pid_val, float) and pid_val.is_integer() else str(pid_val)
                            if allowed_product_ids and pid_str not in allowed_product_ids:
                                skipped_by_store += 1
                                continue
                        else:
                            continue
                        action_in_row = data.get("action_id")
                        if action_in_row not in (None, "", action_id, str(action_id)):
                            continue
                        merged = dict(data)
                        merged["_meta"] = meta or {}
                        candidates.append(merged)
                actions_result.append({
                    "action": action,
                    "candidates": candidates,
                    "headers": sheet_headers,
                    "total": len(candidates),
                })
            if no_snapshot and sheet_rows:
                flash(
                    "Товары не загружены для этого магазина. Обновите товары, чтобы сопоставить Ozon ID из файла.",
                    "warning",
                )
            if skipped_by_store:
                flash(
                    f"{skipped_by_store} строк(и) из файла пропущены: Ozon ID нет в товарах выбранного магазина.",
                    "warning",
                )
        except OzonClientError as exc:
            flash(str(exc), "danger")
        except Exception as exc:
            flash(f"Не удалось получить акции: {exc}", "danger")

    return render_template(
        "actions.html",
        stores=stores,
        selected_store_id=selected_store.id if selected_store else None,
        actions_result=actions_result,
        last_sheet=last_sheet,
    )


@app.route("/api/actions/update", methods=["POST"])
@login_required
def api_actions_update():
    payload = request.get_json(force=True, silent=True) or {}
    store_id = payload.get("store_id")
    action_id = payload.get("action_id")
    mode = payload.get("mode")
    items = payload.get("items", [])

    if not store_id or not action_id or mode not in {"activate", "deactivate"}:
        return {"error": "Некорректные параметры"}, 400

    store = Store.query.get(store_id)
    if not store or (not current_user.is_admin and store.user_id != current_user.id):
        return {"error": "Нет доступа к магазину"}, 403

    client = OzonClient(store.client_id, store.api_key)
    try:
        if mode == "activate":
            products = []
            for item in items:
                pid = item.get("product_id")
                price = item.get("action_price")
                if pid is None or price is None:
                    continue
                products.append({"product_id": pid, "action_price": price})
            if not products:
                return {"error": "Нет товаров для добавления"}, 400
            resp = client.activate_action_products(action_id, products)
        else:
            ids = [item.get("product_id") for item in items if item.get("product_id") is not None]
            if not ids:
                return {"error": "Нет товаров для удаления"}, 400
            resp = client.deactivate_action_products(action_id, ids)
        return {"ok": True, "response": resp}
    except OzonClientError as exc:
        return {"error": str(exc)}, 400
    except Exception as exc:
        return {"error": f"Сбой: {exc}"}, 500


def _prepare_postings(postings, product_lookup, packaging_override=None, promotion_override=None):
    prepared = []
    for posting in postings:
        status = (posting.get("status") or "").lower()
        status_label = STATUS_LABELS.get(status, posting.get("status") or "")
        is_cancelled = status.startswith("cancel")
        fin_products = posting.get("financial_data", {}).get("products", [])
        product_meta_map = {}
        for meta in posting.get("products", []) or []:
            meta_pid = meta.get("product_id") or meta.get("sku")
            if meta_pid is None:
                continue
            product_meta_map[meta_pid] = {
                "offer_id": meta.get("offer_id") or "",
                "name": meta.get("name") or "",
                "price": meta.get("price"),
                "quantity": meta.get("quantity") or 0,
            }

        lines = []
        total_price = 0.0
        payout_sum = 0.0
        commission_sum = 0.0
        discount_sum = 0.0
        total_qty = 0
        expenses_total = 0.0
        cost_total_sum = 0.0
        revenue_sum = 0.0

        # Считаем общее количество единиц по отправлению, чтобы распределить упаковку
        qty_guard = 0
        for product in fin_products:
            qty_guard += int(product.get("quantity", 0) or 0)
        if qty_guard == 0:
            for meta in posting.get("products", []) or []:
                qty_guard += int(meta.get("quantity", 0) or 0)

        packaging_total_posting = (
            packaging_override if packaging_override is not None else PACKAGING_COST
        )

        for product in fin_products:
            product_pid = product.get("product_id") or product.get("sku")
            meta = product_meta_map.get(product_pid, {})
            offer_id = product.get("offer_id") or meta.get("offer_id") or ""
            cached = product_lookup.get(offer_id, {})

            price = float(product.get("price", 0) or meta.get("price") or cached.get("price") or 0)
            quantity = int(product.get("quantity", 0) or meta.get("quantity", 0) or 0)
            payout = float(product.get("payout", 0) or 0)
            cost_value = float(cached.get("cost", 0) or 0)

            commission_raw = float(product.get("commission_amount", 0) or 0)
            commission_percent = float(product.get("commission_percent") or cached.get("commission_percent_fbs") or 0)

            discount_value = float(product.get("total_discount_value", 0) or 0)

            acquiring_unit = float(cached.get("acquiring") or product.get("acquiring") or 0)
            if not acquiring_unit:
                acquiring_unit = price * 0.019
            delivery_unit = float(cached.get("delivery_cost") or 0)
            logistics_unit = float(cached.get("logistics") or 0)
            processing_unit = float(cached.get("processing_cost") or 0)
            promotion_unit = float(
                cached.get("promotion")
                or price * (promotion_override if promotion_override is not None else PROMOTION_RATE)
            )

            price_total = price * quantity
            acquiring_total = acquiring_unit * quantity
            delivery_total = delivery_unit * quantity
            logistics_total = logistics_unit * quantity
            processing_total = processing_unit * quantity
            promotion_total = promotion_unit * quantity

            # Упаковка — одна на отправление, распределяем пропорционально количеству
            packaging_total = 0.0
            if qty_guard:
                packaging_total = packaging_total_posting * (quantity / qty_guard)
            elif fin_products:
                packaging_total = packaging_total_posting / len(fin_products)
            else:
                packaging_total = packaging_total_posting

            # Комиссия: если передана в заказе (обычно отрицательная для доставленных) — берём модуль;
            # если ноль, рассчитываем по проценту sales_percent_fbs от цены * qty
            if commission_raw:
                commission_amount_total = abs(commission_raw)
            else:
                commission_amount_total = price * commission_percent / 100 * quantity if commission_percent else 0

            variable_expenses = (
                acquiring_total
                + delivery_total
                + logistics_total
                + processing_total
                + promotion_total
                + packaging_total
            )

            # Если payout уже есть — он включает комиссию, поэтому повторно не вычитаем
            if payout:
                revenue_total = payout - variable_expenses
                line_expenses = variable_expenses + commission_amount_total
            else:
                line_expenses = variable_expenses + commission_amount_total
                revenue_total = price_total - line_expenses

            cost_total = cost_value * quantity
            margin_value = revenue_total - cost_total
            markup_value = (margin_value / cost_total * 100) if cost_total else 0

            expenses_total += line_expenses
            cost_total_sum += cost_total
            revenue_sum += revenue_total

            total_price += price_total
            payout_sum += payout or revenue_total
            commission_sum += commission_amount_total
            discount_sum += discount_value
            total_qty += quantity

            lines.append(
                {
                    "offer_id": offer_id,
                    "name": product.get("name") or meta.get("name") or cached.get("name") or "",
                    "price": price_total,
                    "quantity": quantity,
                    "payout": payout or revenue_total,
                    "commission": commission_amount_total,
                    "commission_percent": commission_percent,
                    "acquiring": acquiring_total,
                    "delivery_cost": delivery_total,
                    "logistics": logistics_total,
                    "processing": processing_total,
                    "promotion": promotion_total,
                    "packaging": packaging_total,
                    "revenue": revenue_total,
                    "cost": cost_total,
                    "margin": margin_value,
                    "markup": markup_value,
                    "discount": discount_value,
                }
            )

        amount = total_price
        expenses = expenses_total
        revenue = revenue_sum
        margin = revenue - cost_total_sum
        markup = (margin / cost_total_sum * 100) if cost_total_sum else 0

        if status == "cancelled" and not fin_products:
            amount = revenue = margin = expenses = cost_total_sum = 0
            markup = 0
            lines = []

        date_raw = posting.get("in_process_at") or posting.get("created_at") or ""
        date_value = _parse_dt(date_raw, datetime.utcnow()) if date_raw else datetime.utcnow()

        prepared.append(
            {
                "order_number": posting.get("order_number") or posting.get("posting_number") or "",
                "posting_number": posting.get("posting_number") or "",
                "status": posting.get("status") or "",
                "status_label": status_label or posting.get("status") or "",
                "status_key": status or "прочее",
                "is_cancelled": is_cancelled,
                "scheme": posting.get("scheme") or "",
                "date": date_value,
                "amount": total_price,
                "revenue": revenue,
                "margin": margin,
                "markup": markup,
                "cost": cost_total_sum,
                "accruals": payout_sum,
                "items_count": total_qty or len(lines),
                "expenses": expenses,
                "lines": lines,
            }
        )
    prepared.sort(key=lambda x: x.get("date") or datetime.min, reverse=True)
    return prepared


def _markup_color(value: float) -> str:
    try:
        numeric = float(value)
    except (TypeError, ValueError):
        numeric = 0.0
    # 0 % и ниже — насыщенно красный, 40 % и выше — зелёный
    low, high = 0.0, 40.0
    if numeric <= low:
        ratio = 0.0
    elif numeric >= high:
        ratio = 1.0
    else:
        ratio = (numeric - low) / (high - low or 1)
    hue = 120 * ratio  # 0 red -> 120 green
    return f"hsl({hue:.0f}, 75%, 90%)"


def _log_job(job_id: str, message: str, progress: float = None, error: str = None):
    with fetch_lock:
        job = fetch_jobs.setdefault(job_id, {"logs": [], "progress": 0, "status": "running"})
        if progress is not None:
            job["progress"] = progress
        if error:
            job["status"] = "error"
            job["error"] = error
        job["logs"].append({"time": datetime.utcnow().isoformat(), "message": message})
        job["logs"] = job["logs"][-50:]
    logger.info("[%s] %s", job_id, message)
    if error:
        logger.error("[%s] %s", job_id, error)


def _run_fetch_job(job_id: str, store_id: int, user_id: int, packaging: float, promotion_rate: float):
    with app.app_context():
        store = Store.query.filter_by(id=store_id).first()
        if not store:
            _log_job(job_id, "Магазин не найден", error="store_missing")
            return
        # Локальная ссылка, чтобы избежать NameError при многопоточности и явно
        # зафиксировать используемую функцию расчётов в задаче.
        compute_fn = compute_financials
        client = OzonClient(store.client_id, store.api_key)
        error_message = None
        try:
            _log_job(job_id, "Запрашиваем список товаров", progress=2)
            products_raw = []
            last_id = ""
            total_known = 0
            page_limit = 500
            guard = 0
            while True:
                guard += 1
                try:
                    resp = client.list_products(limit=page_limit, last_id=last_id)
                except OzonClientError as exc:
                    error_message = str(exc)
                    _log_job(job_id, f"Не удалось получить часть списка: {exc}")
                    logger.exception("Ошибка list_products для job %s", job_id)
                    break
                page_items = resp.get("result", {}).get("items", [])
                products_raw.extend(page_items)
                total_known = resp.get("result", {}).get("total") or total_known or len(page_items)
                last_id = resp.get("result", {}).get("last_id") or ""
                done_ratio = min(0.2, (len(products_raw) / max(total_known, 1)) * 0.2)
                _log_job(job_id, f"Получено товаров: {len(products_raw)}", progress=5 + done_ratio * 100)
                if not last_id or not page_items or guard > 200:
                    break

            offer_ids = [item.get("offer_id") for item in products_raw if item.get("offer_id")]
            avg_delivery = None
            try:
                avg_delivery = client.average_delivery_time()
            except Exception:
                avg_delivery = None

            if error_message and not offer_ids:
                _log_job(job_id, "Данных нет из-за ошибки запросов", error=error_message)
                return

            _log_job(job_id, "Рассчитываем показатели", progress=50)

            temp_file = tempfile.NamedTemporaryFile(
                mode="w+", delete=False, encoding="utf-8", suffix=".jsonl"
            )
            total_items = 0
            failed_calc = 0
            failed_info = 0
            failed_price = 0
            preview: list = []
            cache_rows: list = []

            chunks = list(client._chunked(offer_ids, 500)) if offer_ids else []
            for idx, chunk in enumerate(chunks, 1):
                info_map = {}
                price_map = {}
                try:
                    info_resp = client.product_info(chunk)
                    for item in info_resp.get("items", []):
                        info_map[item.get("offer_id")] = item
                except OzonClientError as exc:
                    failed_info += len(chunk)
                    _log_job(job_id, f"Инфо {idx}/{len(chunks)} пропущено: {exc}")
                    logger.exception("Ошибка info chunk %s/%s job %s", idx, len(chunks), job_id)

                cursor = ""
                try:
                    while True:
                        price_resp = client.product_prices(chunk, limit=500, cursor=cursor)
                        for item in price_resp.get("items", []):
                            price_map[item.get("offer_id")] = item
                        cursor = price_resp.get("cursor") or ""
                        if not cursor:
                            break
                except OzonClientError as exc:
                    failed_price += len(chunk)
                    _log_job(job_id, f"Цены {idx}/{len(chunks)} пропущено: {exc}")
                    logger.exception("Ошибка price chunk %s/%s job %s", idx, len(chunks), job_id)

                for offer_id in chunk:
                    prepared_input = {
                        "offer_id": offer_id,
                        "info": info_map.get(offer_id, {}),
                        "price": price_map.get(offer_id, {}),
                        "delivery": avg_delivery or {},
                    }
                    try:
                        prepared = compute_fn(prepared_input, packaging, promotion_rate)
                        if len(preview) < 500:
                            preview.append(prepared)
                        cache_rows.append(prepared)
                        temp_file.write(json.dumps(prepared, ensure_ascii=False) + "\n")
                        total_items += 1
                    except Exception as exc:
                        failed_calc += 1
                        logger.exception("Ошибка расчёта для %s в job %s", offer_id, job_id)
                        _log_job(job_id, f"Не удалось посчитать {offer_id}: {exc}")

                _log_job(job_id, f"Инфо {idx}/{len(chunks)}", progress=20 + (idx / max(len(chunks), 1)) * 25)
                _log_job(job_id, f"Цены {idx}/{len(chunks)}", progress=45 + (idx / max(len(chunks), 1)) * 35)

            temp_file.flush()
            temp_file.close()

            if failed_info:
                _log_job(job_id, f"Пропущено из-за ошибок инфо: {failed_info}")
            if failed_price:
                _log_job(job_id, f"Пропущено из-за ошибок цен: {failed_price}")
            if failed_calc:
                _log_job(job_id, f"Пропущено из-за ошибок расчёта: {failed_calc}")

            # На всякий случай синхронизируем счётчик с реальным количеством записей
            total_items = total_items or len(cache_rows)

            snapshot = ProductSnapshot(store_id=store.id, user_id=user_id, data=cache_rows)
            db.session.add(snapshot)
            db.session.commit()

            _log_job(job_id, "Готово", progress=100)

            with fetch_lock:
                fetch_jobs[job_id] = {
                    "status": "done",
                    "progress": 100,
                    "logs": fetch_jobs.get(job_id, {}).get("logs", []),
                    "result_path": temp_file.name,
                    "total": total_items,
                    "error": error_message,
                }
        except Exception as exc:
            trace = traceback.format_exc()
            _log_job(job_id, f"Ошибка загрузки: {exc.__class__.__name__}", error=str(exc))
            logger.error("Фоновая задача %s упала: %s\n%s", job_id, exc, trace)
            with fetch_lock:
                job = fetch_jobs.get(job_id, {})
                job["status"] = "error"
                job["error"] = str(exc)
                job["trace"] = trace
                fetch_jobs[job_id] = job


def _start_fetch_job(store_id: int, user_id: int, packaging: float, promotion_rate: float) -> str:
    job_id = uuid.uuid4().hex
    _log_job(job_id, "Запуск задачи", progress=1)
    thread = threading.Thread(
        target=_run_fetch_job,
        args=(job_id, store_id, user_id, packaging, promotion_rate),
        daemon=True,
    )
    thread.start()
    return job_id


@app.route("/products", methods=["GET", "POST"])
@login_required
def products():
    stores = Store.query.all() if current_user.is_admin else Store.query.filter_by(user_id=current_user.id).all()
    selected_store_id = request.values.get("store_id", type=int)
    packaging_override = request.values.get("packaging", type=float)
    promotion_override = request.values.get("promotion", type=float)
    promotion_rate = promotion_override / 100 if promotion_override is not None else None
    table_data = []
    raw_products = []
    if selected_store_id:
        snapshot = (
            ProductSnapshot.query.filter_by(store_id=selected_store_id)
            .order_by(ProductSnapshot.created_at.desc())
            .first()
        )
        if snapshot:
            raw_products = snapshot.data
            table_data = prepare_table_data(raw_products, packaging_override, promotion_rate)
    if request.method == "POST" and request.form.get("action") == "fetch":
        if current_user.is_limited and not current_user.is_admin:
            flash("Ваш доступ к обновлению данных ограничен администратором", "warning")
            return redirect(url_for("products"))
        store = Store.query.filter_by(id=selected_store_id).first()
        if not current_user.is_admin and store and store.user_id != current_user.id:
            flash("Нет доступа к магазину", "danger")
            return redirect(url_for("products"))
        if not store:
            flash("Выберите магазин", "warning")
            return redirect(url_for("products"))
        try:
            client = OzonClient(store.client_id, store.api_key)
            raw_products = client.fetch_joined_products()
            table_data = prepare_table_data(raw_products, packaging_override, promotion_rate)
            snapshot = ProductSnapshot(store_id=store.id, user_id=store.user_id, data=table_data)
            db.session.add(snapshot)
            db.session.commit()
            flash("Данные обновлены", "success")
        except OzonClientError as exc:
            flash(str(exc), "danger")
        except Exception as exc:
            flash(f"Не удалось получить данные: {exc}", "danger")
    return render_template(
        "products.html",
        stores=stores,
        selected_store_id=selected_store_id,
        table_data=table_data,
        raw=json.dumps(table_data),
        packaging_value=packaging_override if packaging_override is not None else PACKAGING_COST,
        promotion_value=promotion_override if promotion_override is not None else PROMOTION_RATE * 100,
    )


@app.route("/api/products/start_fetch", methods=["POST"])
@login_required
def start_fetch():
    data = request.get_json(silent=True) or {}
    store_id = data.get("store_id")
    try:
        packaging = float(data.get("packaging")) if data.get("packaging") is not None else PACKAGING_COST
    except (TypeError, ValueError):
        packaging = PACKAGING_COST
    try:
        promotion = float(data.get("promotion")) if data.get("promotion") is not None else PROMOTION_RATE * 100
    except (TypeError, ValueError):
        promotion = PROMOTION_RATE * 100
    promotion_rate = promotion / 100

    store = Store.query.filter_by(id=store_id).first()
    if not store:
        return {"ok": False, "error": "Магазин не найден"}, 404
    if not current_user.is_admin and store.user_id != current_user.id:
        return {"ok": False, "error": "Нет доступа к магазину"}, 403

    job_id = _start_fetch_job(store.id, current_user.id, packaging, promotion_rate)
    return {"ok": True, "job_id": job_id}


@app.route("/api/products/progress/<job_id>", methods=["GET"])
@login_required
def fetch_progress(job_id):
    with fetch_lock:
        job = fetch_jobs.get(job_id)
    if not job:
        # Возвращаем 200, чтобы фронт не падал по fail() и отобразил ошибку
        return {"ok": False, "status": "missing", "error": "Задача не найдена"}
    response = {
        "ok": True,
        "status": job.get("status", "running"),
        "progress": job.get("progress", 0),
        "logs": job.get("logs", []),
    }
    if job.get("status") == "done":
        response["total"] = job.get("total")
    if job.get("error"):
        response["error"] = job.get("error")
    if job.get("trace"):
        response["trace"] = job.get("trace")
    return response


@app.route("/api/products/result/<job_id>", methods=["GET"])
@login_required
def fetch_result(job_id):
    try:
        limit = int(request.args.get("limit", 1000))
    except ValueError:
        limit = 1000
    try:
        offset = int(request.args.get("offset", 0))
    except ValueError:
        offset = 0
    limit = max(1, min(limit, 2000))
    offset = max(0, offset)

    with fetch_lock:
        job = fetch_jobs.get(job_id)
    if not job:
        return {"ok": False, "error": "Задача не найдена"}, 404
    if job.get("status") != "done":
        return {"ok": False, "error": "Задача ещё выполняется"}, 400

    total = int(job.get("total", 0))
    items = []
    has_more = False

    result_path = job.get("result_path")
    if result_path and os.path.exists(result_path):
        try:
            if not total:
                try:
                    with open(result_path, "r", encoding="utf-8") as fh:
                        total = sum(1 for _ in fh)
                    with fetch_lock:
                        job["total"] = total
                        fetch_jobs[job_id] = job
                except OSError:
                    total = 0
            with open(result_path, "r", encoding="utf-8") as fh:
                for line in itertools.islice(fh, offset, offset + limit):
                    try:
                        items.append(json.loads(line))
                    except json.JSONDecodeError:
                        continue
            has_more = offset + limit < total if total else False
        except OSError as exc:
            return {"ok": False, "error": f"Не удалось прочитать результаты: {exc}"}, 500
    else:
        items = []

    return {
        "ok": True,
        "items": items,
        "total": total,
        "next_offset": offset + limit if has_more else None,
        "has_more": has_more,
    }


@app.route("/products/update_price", methods=["POST"])
@login_required
def update_price():
    if current_user.is_limited and not current_user.is_admin:
        return {"ok": False, "error": "Ваш доступ к изменению цен ограничен"}, 403

    data = request.get_json(silent=True) or {}
    store_id = data.get("store_id")
    items = data.get("items") or []
    if not store_id or not items:
        return {"ok": False, "error": "Нет данных для обновления"}, 400

    store = Store.query.filter_by(id=store_id).first()
    if not store:
        return {"ok": False, "error": "Магазин не найден"}, 404
    if not current_user.is_admin and store.user_id != current_user.id:
        return {"ok": False, "error": "Нет доступа к магазину"}, 403

    prepared_prices = []
    for item in items:
        offer_id = item.get("offer_id")
        try:
            price_value = round(float(item.get("price") or 0), 2)
            cost_value = round(float(item.get("cost") or 0), 2)
            min_price_value = round(float(item.get("min_price") or price_value), 2)
        except (TypeError, ValueError):
            continue
        if not offer_id or price_value <= 0:
            continue
        prepared_prices.append(
            {
                "offer_id": offer_id,
                "price": f"{price_value:.2f}",
                "old_price": "0",
                "min_price": f"{min_price_value:.2f}",
                "net_price": f"{cost_value:.2f}",
                "currency_code": "RUB",
                "auto_action_enabled": "ENABLED",
                "auto_add_to_ozon_actions_list_enabled": "DISABLED",
                "manage_elastic_boosting_through_price": True,
            }
        )

    if not prepared_prices:
        return {"ok": False, "error": "Нет валидных товаров для обновления"}, 400

    try:
        client = OzonClient(store.client_id, store.api_key)
        response = client.import_prices(prepared_prices)
        return {"ok": True, "result": response}
    except OzonClientError as exc:
        return {"ok": False, "error": str(exc)}, 400
    except Exception as exc:
        return {"ok": False, "error": f"Ошибка обновления: {exc}"}, 500


@app.template_filter("datetime")
def format_datetime(value):
    if isinstance(value, datetime):
        return value.strftime("%d.%m.%Y %H:%M")
    return value


@app.template_filter("markup_color")
def markup_color_filter(value):
    return _markup_color(value)


if __name__ == "__main__":
    app.run(debug=True, host="0.0.0.0", port=5000)
